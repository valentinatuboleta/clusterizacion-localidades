"""
Tests herméticos para la selección estratificada de tarjetas de presentación
y el anexo de marcha blanca (Bloques 1 y 3).
"""

import os
import pytest
import pandas as pd
import openpyxl

from scripts.preparar_ejemplos_presentacion import (
    preparar_galeria_extendida,
    generar_anexo_excel
)

CSV_GALERIA_PATH = "reports/ejemplos_galeria_extendida.csv"
EXCEL_ANEXO_PATH = "reports/anexo_asignaciones_marcha_blanca.xlsx"
CSV_ASIGNACIONES_PATH = "data/processed/asignacion_microclusters.csv"


def test_existencia_artefactos():
    """Valida que los artefactos generados existan en disco."""
    assert os.path.exists(CSV_GALERIA_PATH), f"No existe {CSV_GALERIA_PATH}"
    assert os.path.exists(EXCEL_ANEXO_PATH), f"No existe {EXCEL_ANEXO_PATH}"


def test_columnas_galeria_extendida():
    """Valida el esquema de columnas del CSV de galería extendida."""
    df = pd.read_csv(CSV_GALERIA_PATH)
    columnas_esperadas = {
        "localidad", "recinto", "precio", "type_site", "micro_cluster_id",
        "label_auto", "arquetipo_demanda", "score_confianza", "es_frontera",
        "grupo", "caso"
    }
    assert columnas_esperadas.issubset(set(df.columns)), (
        f"Faltan columnas en {CSV_GALERIA_PATH}. Esperadas: {columnas_esperadas}"
    )


def test_conteos_por_grupo():
    """
    Valida que los conteos respeten los límites especificados:
      - Premium: VIP + Preferencial (<= 8 tarjetas)
      - Masivos: Platea General + Popular + Grada (<= 12 tarjetas)
      - Especiales: AU-0 (2) + 2 fronteras + 1 híbrido (= 5 tarjetas)
    """
    df = pd.read_csv(CSV_GALERIA_PATH)
    
    n_premium = len(df[df["grupo"] == "premium"])
    n_masivos = len(df[df["grupo"] == "masivos"])
    n_especiales = len(df[df["grupo"] == "especiales"])
    
    assert n_premium <= 8, f"Grupo premium excede tope de 8: tiene {n_premium}"
    assert n_masivos <= 12, f"Grupo masivos excede tope de 12: tiene {n_masivos}"
    assert n_especiales == 5, f"Grupo especiales debe tener exactamente 5: tiene {n_especiales}"


def test_ausencia_clusters_degenerados():
    """
    VIP-3 y PPF-3 deben estar ausentes por la compuerta de degeneración
    en toda la galería extendida.
    """
    df = pd.read_csv(CSV_GALERIA_PATH)
    clusters_presentes = set(df["micro_cluster_id"].dropna().unique())
    
    assert "VIP-3" not in clusters_presentes, "VIP-3 está presente pero debió ser excluido"
    assert "PPF-3" not in clusters_presentes, "PPF-3 está presente pero debió ser excluido"


def test_valores_y_consistencia_columna_caso():
    """
    Valida los valores de la columna caso:
      - Solo admite: representante, frontera, hibrido.
      - Filas marcadas frontera deben tener es_frontera == True.
      - Filas marcadas representante deben tener es_frontera == False.
      - Filas marcadas hibrido deben tener label_auto empezando por 'hibrido'.
    """
    df = pd.read_csv(CSV_GALERIA_PATH)
    valores_permitidos = {"representante", "frontera", "hibrido"}
    valores_encontrados = set(df["caso"].unique())
    
    assert valores_encontrados.issubset(valores_permitidos), (
        f"Valores no permitidos en columna caso: {valores_encontrados - valores_permitidos}"
    )
    
    # Consistencia frontera
    df_fronteras = df[df["caso"] == "frontera"]
    assert len(df_fronteras) == 2, f"Deben ser exactamente 2 fronteras: hay {len(df_fronteras)}"
    assert (df_fronteras["es_frontera"] == True).all(), (
        "Todas las filas marcadas como frontera deben tener es_frontera == True"
    )
    
    # Consistencia representante
    df_reps = df[df["caso"] == "representante"]
    assert (df_reps["es_frontera"] == False).all(), (
        "Los representantes no deben ser localidades de frontera"
    )
    
    # Consistencia híbrido
    df_hibridos = df[df["caso"] == "hibrido"]
    assert len(df_hibridos) == 1, f"Debe haber exactamente 1 híbrido: hay {len(df_hibridos)}"
    assert df_hibridos.iloc[0]["label_auto"].startswith("hibrido"), (
        f"El híbrido debe tener label_auto comenzando por 'hibrido': {df_hibridos.iloc[0]['label_auto']}"
    )


def test_anexo_excel_integridad():
    """
    Valida que el libro Excel generado cumpla con los requisitos:
      - Hoja única 'Asignaciones_Marcha_Blanca'
      - Exactamente 1,324 filas de datos + encabezado
      - Autofiltro activado
      - Congelación de panel (fila 1 congelada)
      - Columnas requeridas presentes
    """
    wb = openpyxl.load_workbook(EXCEL_ANEXO_PATH)
    
    assert len(wb.sheetnames) == 1, f"Debe tener exactamente una hoja: {wb.sheetnames}"
    ws = wb.active
    assert ws.title == "Asignaciones_Marcha_Blanca"
    
    # 1 encabezado + 1324 filas
    assert ws.max_row == 1325, f"Debe tener 1,325 filas (1 + 1324): tiene {ws.max_row}"
    
    # Autofiltro activado
    assert ws.auto_filter.ref is not None, "El autofiltro no está configurado en la hoja"
    
    # Fila 1 congelada
    assert ws.freeze_panes == "A2", f"La fila 1 debe estar congelada (A2): tiene {ws.freeze_panes}"
    
    # Columnas esperadas en encabezado
    encabezados = [cell.value for cell in ws[1]]
    columnas_esperadas = [
        "localidad", "recinto", "precio", "type_site", "micro_cluster_id",
        "label_auto", "arquetipo_demanda", "score_confianza", "es_frontera"
    ]
    assert encabezados == columnas_esperadas, (
        f"Encabezados no coinciden: {encabezados} vs {columnas_esperadas}"
    )
